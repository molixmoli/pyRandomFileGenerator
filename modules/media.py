import os
from tkinter import Image
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont
from docx import Document
import cv2

class Media:

    def generate_big_random_bin_file(self, filename, size):
        """
        generate big binary file with the specified size in bytes
        :param filename: the filename
        :param size: the size in bytes
        :return:void
        """
        import os
        with open('%s' % filename, 'wb') as fout:
            fout.write(os.urandom(size))  # 1
        #print('big random binary file with size %f generated ok' % size)
        pass


    def generate_big_random_letters(self, filename, size):
        """
        generate big random letters/alphabets to a file
        :param filename: the filename
        :param size: the size in bytes
        :return: void
        """
        import random
        import string

        chars = ''.join([random.choice(string.ascii_letters) for i in range(size)])  # 1

        with open(filename, 'w') as f:
            f.write(chars)
        pass


    def generate_big_sparse_file(self, filename, size):
        f = open(filename, "wb")
        f.seek(size - 1)
        f.write(b"\1")
        f.close()
        pass

    def generate_big_random_sentences(self, filename, linecount):
        import random

        reg1 = ("Mami", "Gata", "Perra", "Zorra", "Chica")
        reg2 = ("yo quiero", "vamos a", "yo voy a", "yo quiero", "yo vengo a")
        reg3 = ("castigarte", "cogerte", "encenderte", "darte", "azotarte")
        reg4 = ("duro", "rápido", "lento", "suave", "fuerte")
        reg5 = ("hasta que salga el sol", "toda la noche", "hasta el amanecer", "hasta mañana", "todo el día")
        reg6 = ("sin miedo", "sin anestesia", "en el piso", "contra la pared", "sin compromiso")

        #all = [reg1, reg2, reg3, reg4, reg5, reg6]

        col1 = ("Queridos compañeros", "Por otra parte,y dados los condicionamientos actuales", "Asimismo,",
                "Sin embargo no hemos de olvidar que", "De igual manera,", "La práctica de la vida cotidiana prueba que,",
                "No es indispensable argumentar el peso y la significación de estos problemas ya que,",
                "Las experiencias ricas y diversas muestran que,", "El afán de organización, pero sobre todo",
                "Los superiores principios ideológicos, condicionan que", "Incluso, bien pudiéramos atrevernos a sugerir que",
                "Es obvio señalar que,", "Pero pecaríamos de insinceros si soslayásemos que,",
                "Y además, quedaríamos inmersos en la más abyecta de las estulticias si no fueramos consacientes de que,",
                "Por último, y como definitivo elemento esclarecedor, cabe añadir que,")
        col2 = ("la realización de las premisas del programa", "la complejidad de los estudios de los dirigentes",
                "el aumento constante, en cantidad y en extensión, de nuestra actividad", "la estructura actual de la organización",
                "el nuevo modelo de actividad de la organización,", "el desarrollo continuo de distintas formas de actividad",
                "nuestra actividad de información y propaganda", "el reforzamiento y desarrollo de las estructuras",
                "la consulta con los numerosos militantes", "el inicio de la acción general de formación de las actitudes",
                "un relanzamiento específico de todos los sectores implicados", "la superación de experiencias periclitadas",
                "una aplicación indiscriminada de los factores confluyentes", "la condición sine qua non rectora del proceso",
                "el proceso consensuado de unas y otras aplicaciones concurrentes")
        col3 = ("nos obliga a un exhaustivo análisis", "cumple un rol escencial en la formación",
                "exige la precisión y la determinación", "ayuda a la preparación y a la realización",
                "garantiza la participación de un grupo importante en la formación", "cumple deberes importantes en la determinación",
                "facilita la creación", "obstaculiza la apreciación de la importancia", "ofrece un ensayo interesante de verificación",
                "implica el proceso de reestructuración y modernización", "habrá de significar un auténtico y eficaz punto de partida",
                "permite en todo caso explicitar las razones fundamentales", "asegura, en todo caso, un proceso muy sensible de inversión",
                "radica en una elaboración cuidadosa y sistemática de las estrategias adecuadas", "deriva de una indirecta incidencia superadora")
        col4 = ("de las condiciones financieras y administrativas existentes.", "de las directivas de desarrollo para el futuro.",
                "del sistema de participación general.", "de las actitudes de los miembros hacia sus deberes ineludibles.",
                "de las nuevas proposiciones.", "de las direcciones educativas en el sentido del progreso.",
                "del sistema de formación de cuadros que corresponda a las necesidades.", "de las condiciones de las actividades apropiadas.",
                "del modelo de desarrollo.", "de las formas de acción.", "de las básicas premisas adoptadas.",
                "de toda una casuística de amplio espectro.", "de los elementos generadores.",
                "para configurar una interface amigable y coadyuvante a la reingeniería del sistema.",
                "de toda una serie de criterios ideológicamente sistematizados en un frente común de actuación regeneradora.")

        all = [col1, col2, col3, col4]

        with open(filename, 'w') as f:
            for i in range(linecount):
                f.writelines([' '.join([random.choice(i) for i in all]), '\n'])
        pass


    def generate_pdf(self, filename_input, filename_output):
        # save FPDF() class into
        # a variable pdf
        pdf = FPDF()

        # Add a page
        pdf.add_page()

        # set style and size of font
        # that you want in the pdf
        pdf.set_font("Arial", size=10)

        # open the text file in read mode
        f = open(filename_input, "r")

        # insert the texts in pdf
        for x in f:
            pdf.cell(200, 10, txt=x, ln=1, align='C')

        # save the pdf with name .pdf
        pdf.output(filename_output)
        pass


    def generate_img(self, filename_input, filename_output):
        width = 1024
        height = 512
        message = ''

        img = Image.open("./resources/watermark.png").resize((width, height))
        imgDraw = ImageDraw.Draw(img)

        f = open(filename_input, "r")
        for x in f:
            message = message + x
        f.close()

        font = ImageFont.truetype("arial.ttf", size=20)

        imgDraw.text((10, 10), message, font=font, fill=(255, 255, 0))
        img = img.convert('RGB')
        img.save(filename_output)
        pass

    def generate_docx(self, filename_input, filename_output):
        doc = Document()

        with open(filename_input, 'r', encoding='latin-1') as openfile:
            line = openfile.read()
            doc.add_paragraph(line)
            doc.save(filename_output)
        pass

    def generate_pdf_from_docx(self, filename_input, filename_output):
        import win32com.client

        wdFormatPDF = 17

        inputFile = os.path.abspath(filename_input)
        outputFile = os.path.abspath(filename_output)
        word = win32com.client.Dispatch('Word.Application')
        doc = word.Documents.Open(inputFile)
        doc.SaveAs(outputFile, FileFormat=wdFormatPDF)
        doc.Close()
        word.Quit()
        pass

    def generate_avi(self, filename_avi, directory):
        images = [img for img in os.listdir(directory) if img.endswith(".png")]
        frame = cv2.imread(os.path.join(directory, images[0]))
        height, width, layers = frame.shape

        video = cv2.VideoWriter(filename_avi, 0, 1, (width,height))

        for image in images:
            video.write(cv2.imread(os.path.join(directory, image)))

        cv2.destroyAllWindows()
        video.release()
        pass

    def generate_audio(self, filename_input, filename_output):
        from gtts import gTTS
        from pydub import AudioSegment

        file = open(filename_input, "r").read().replace("\n", " ")

        language = "es"
        region = "es"
        speech = gTTS(text=str(file), lang=language, tld=region, slow=True)

        speech.save(filename_output + ".mp3")
        # playsound(audio_file)

        sound1 = AudioSegment.from_file("./resources/base.mp3", format="mp3")
        sound2 = AudioSegment.from_file(filename_output + ".mp3", format="mp3")

        # Overlay sound2 over sound1 at position 0  (use louder instead of sound1 to use the louder version)
        # overlay = sound1.overlay(sound2, position=0)
        overlay = sound1.overlay(sound2, position=0.1 * len(sound1))

        # simple export
        file_handle = overlay.export(filename_output + " hit.mp3", format="mp3")

        pass